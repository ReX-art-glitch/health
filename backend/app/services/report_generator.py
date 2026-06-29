from typing import Dict, List, Any
from datetime import datetime
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fpdf import FPDF
import io
import openai

from ..config import settings

class ReportGenerator:
    """Generate comprehensive public health reports"""
    
    def __init__(self):
        self.openai_client = openai.OpenAI(api_key=settings.OPENAI_API_KEY)
    
    async def generate_weekly_report(self, data: Dict) -> bytes:
        """Generate weekly public health report"""
        doc = Document()
        
        # Title
        title = doc.add_heading('Weekly Public Health Report', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Date
        doc.add_paragraph(f'Report Date: {datetime.now().strftime("%B %d, %Y")}')
        doc.add_paragraph(f'Week: {datetime.now().isocalendar()[1]}')
        doc.add_paragraph('')
        
        # Executive Summary
        doc.add_heading('Executive Summary', level=1)
        summary = await self.generate_executive_summary(data)
        doc.add_paragraph(summary)
        
        # Vaccination Coverage
        doc.add_heading('Vaccination Coverage', level=1)
        if 'vaccination' in data:
            self._add_vaccination_section(doc, data['vaccination'])
        
        # Disease Surveillance
        doc.add_heading('Disease Surveillance', level=1)
        if 'diseases' in data:
            self._add_disease_section(doc, data['diseases'])
        
        # Maternal Health
        doc.add_heading('Maternal Health', level=1)
        if 'maternal' in data:
            self._add_maternal_section(doc, data['maternal'])
        
        # Drug Inventory
        doc.add_heading('Drug Inventory Status', level=1)
        if 'inventory' in data:
            self._add_inventory_section(doc, data['inventory'])
        
        # Recommendations
        doc.add_heading('AI-Generated Recommendations', level=1)
        if 'recommendations' in data:
            for rec in data['recommendations']:
                doc.add_paragraph(f'• {rec}')
        
        # Save to bytes
        doc_bytes = io.BytesIO()
        doc.save(doc_bytes)
        doc_bytes.seek(0)
        
        return doc_bytes.getvalue()
    
    async def generate_executive_summary(self, data: Dict) -> str:
        """Generate AI executive summary"""
        prompt = f"""
        Create a concise executive summary for a public health report with this data:
        {data}
        
        Include:
        1. Key achievements
        2. Critical challenges
        3. Areas requiring immediate attention
        4. Overall assessment
        
        Keep it under 300 words, professional and actionable.
        """
        
        try:
            response = self.openai_client.chat.completions.create(
                model="gpt-4",
                messages=[{"role": "user", "content": prompt}],
                max_tokens=500
            )
            return response.choices[0].message.content
        except:
            return "Executive summary generation failed."
    
    def _add_vaccination_section(self, doc: Document, data: Dict):
        """Add vaccination coverage details"""
        doc.add_paragraph(f"Overall Coverage: {data.get('overall_coverage', 'N/A')}%")
        
        if 'coverage_by_location' in data:
            table = doc.add_table(rows=1, cols=2, style='Table Grid')
            header = table.rows[0].cells
            header[0].text = 'Location'
            header[1].text = 'Coverage %'
            
            for location, coverage in data['coverage_by_location'].items():
                row = table.add_row()
                row.cells[0].text = location
                row.cells[1].text = f"{coverage:.1f}%"
    
    def _add_disease_section(self, doc: Document, data: Dict):
        """Add disease surveillance details"""
        for disease, info in data.items():
            doc.add_paragraph(f"{disease}:", style='List Bullet')
            doc.add_paragraph(f"  Cases: {info.get('cases', 0)}")
            doc.add_paragraph(f"  Trend: {info.get('trend', 'stable')}")
    
    def _add_maternal_section(self, doc: Document, data: Dict):
        """Add maternal health details"""
        doc.add_paragraph(f"ANC Coverage: {data.get('anc_coverage', 'N/A')}%")
        doc.add_paragraph(f"Deliveries: {data.get('deliveries', 0)}")
        doc.add_paragraph(f"High-Risk Cases: {data.get('high_risk', 0)}")
    
    def _add_inventory_section(self, doc: Document, data: Dict):
        """Add drug inventory details"""
        table = doc.add_table(rows=1, cols=3, style='Table Grid')
        header = table.rows[0].cells
        header[0].text = 'Drug'
        header[1].text = 'Stock'
        header[2].text = 'Status'
        
        for drug, info in data.items():
            row = table.add_row()
            row.cells[0].text = drug
            row.cells[1].text = str(info.get('quantity', 0))
            row.cells[2].text = 'LOW' if info.get('quantity', 0) < info.get('threshold', 100) else 'OK'
    
    async def generate_pdf_report(self, data: Dict) -> bytes:
        """Generate PDF report"""
        pdf = FPDF()
        pdf.add_page()
        
        # Title
        pdf.set_font('Arial', 'B', 16)
        pdf.cell(0, 10, 'Public Health Report', 0, 1, 'C')
        pdf.set_font('Arial', '', 12)
        pdf.cell(0, 10, f'Date: {datetime.now().strftime("%B %d, %Y")}', 0, 1, 'C')
        pdf.ln(10)
        
        # Add sections (simplified for PDF)
        pdf.set_font('Arial', 'B', 14)
        pdf.cell(0, 10, 'Executive Summary', 0, 1)
        pdf.set_font('Arial', '', 11)
        summary = await self.generate_executive_summary(data)
        pdf.multi_cell(0, 10, summary)
        
        # Output
        return pdf.output(dest='S').encode('latin1')